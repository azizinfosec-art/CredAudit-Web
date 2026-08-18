import os

TEXT_EXTS={'.txt','.json','.env','.log','.cfg','.ini','.yaml','.yml','.py','.js','.toml'}
KEYWORDS = {"password","pass","pwd","secret","apikey","api_key","api-key","token"}
USERNAME_KEYWORDS = {"username","user_id","userid","login","user","email"}
HEADER_KEYWORDS = KEYWORDS | USERNAME_KEYWORDS
HEADER_TEXT_CHARS = set(" _-/().")


def _normalize_cell_label(value):
    label = str(value or "").strip().lower()
    label = label.strip(":=-> \t\r\n\"'`")
    label = label.replace(" ", "_").replace("-", "_")
    if label == "api_key":
        return "api_key"
    return label


def _cell_text(value):
    text = str(value).strip() if value is not None else ""
    return " ".join(text.split())


def _looks_like_header_label(value):
    return _normalize_cell_label(value) in HEADER_KEYWORDS


def _looks_like_table_header_cell(value):
    if _looks_like_header_label(value):
        return True
    text = _cell_text(value)
    if not (1 <= len(text) <= 40):
        return False
    if not any(ch.isalpha() for ch in text):
        return False
    if any(ch.isdigit() for ch in text):
        return False
    return all(ch.isalpha() or ch in HEADER_TEXT_CHARS for ch in text)


def _looks_like_table_header_row(cells):
    values = [value for _col, _coord, _row_num, value in cells]
    if not any(_looks_like_header_label(value) for value in values):
        return False
    if all(_looks_like_header_label(value) for value in values):
        return True
    return len(values) >= 3 and all(_looks_like_table_header_cell(value) for value in values)


def _extract_xlsx_text(p):
    import openpyxl
    import warnings

    out = []
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message=r"Data Validation extension is not supported.*",
            category=UserWarning,
            module=r"openpyxl\..*",
        )
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            for sheet_index, ws in enumerate(wb.worksheets, start=1):
                prev_by_col = {}
                for row in ws.iter_rows():
                    cells = []
                    for cell in row:
                        value = _cell_text(cell.value)
                        if value:
                            cells.append((cell.column, cell.coordinate, cell.row, value))
                    if not cells:
                        continue

                    current_by_col = {col: value for col, _coord, _row_num, value in cells}
                    if _looks_like_table_header_row(cells):
                        prev_by_col = current_by_col
                        continue

                    parts = []
                    used_cols = set()

                    # Common table layout: headers in one row, values in the next row.
                    for col, _coord, _row_num, value in cells:
                        header = prev_by_col.get(col)
                        if header and _looks_like_header_label(header) and not _looks_like_header_label(value):
                            parts.append(f"{header.rstrip(':')}: {value}")
                            used_cols.add(col)

                    # Common key/value layout: password | Secret123 in the same row.
                    i = 0
                    while i < len(cells) - 1:
                        col, _coord, _row_num, value = cells[i]
                        next_col, _next_coord, _next_row_num, next_value = cells[i + 1]
                        if (
                            col not in used_cols
                            and next_col not in used_cols
                            and next_col == col + 1
                            and _looks_like_header_label(value)
                            and not _looks_like_header_label(next_value)
                        ):
                            parts.append(f"{value.rstrip(':')}: {next_value}")
                            used_cols.update({col, next_col})
                            i += 2
                            continue
                        i += 1

                    for col, coord, _row_num, value in cells:
                        if col not in used_cols:
                            parts.append(f"cell {coord} {value}")

                    if parts:
                        out.append(f"worksheet {sheet_index} row {cells[0][2]} " + " ".join(parts))
        finally:
            try:
                wb.close()
            except Exception:
                pass
    return "\n".join(out)


def _text_decode_score(text):
    if not text:
        return -1000.0
    length = len(text)
    ascii_like = 0
    printable = 0
    bad_controls = 0
    for ch in text:
        code = ord(ch)
        if ch in "\r\n\t" or 32 <= code <= 126:
            ascii_like += 1
        if ch in "\r\n\t" or ch.isprintable():
            printable += 1
        if code < 32 and ch not in "\r\n\t":
            bad_controls += 1
    nul_count = text.count("\x00")
    replacement_count = text.count("\ufffd")
    return (
        (ascii_like / length) * 2.0
        + (printable / length)
        - (nul_count / length) * 8.0
        - (replacement_count / length) * 6.0
        - (bad_controls / length) * 4.0
    )


def read_text_with_fallback(p):
    try:
        with open(p, "rb") as f:
            data = f.read()
    except Exception:
        return None
    if not data:
        return ""

    bom_encodings = [
        (b"\xef\xbb\xbf", "utf-8-sig"),
        (b"\xff\xfe", "utf-16"),
        (b"\xfe\xff", "utf-16"),
    ]
    for bom, enc in bom_encodings:
        if data.startswith(bom):
            try:
                return data.decode(enc)
            except Exception:
                break

    if b"\x00" not in data[:4096]:
        try:
            return data.decode("utf-8")
        except UnicodeDecodeError:
            pass

    candidates = []
    for enc in ("utf-8", "utf-16", "utf-16-le", "utf-16-be", "latin-1"):
        try:
            text = data.decode(enc)
        except Exception:
            continue
        candidates.append((_text_decode_score(text), text))
    if not candidates:
        return None
    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0][1]
def extract_text_from_file(p):
    ext=os.path.splitext(p)[1].lower()
    try:
        if ext in TEXT_EXTS: 
            return read_text_with_fallback(p)
        if ext=='.pdf':
            try:
                from pdfminer.high_level import extract_text
                return extract_text(p)
            except Exception:
                return None
        if ext=='.docx':
            try:
                from docx import Document
                return "\n".join([x.text for x in Document(p).paragraphs])
            except Exception:
                return None
        if ext=='.xlsx':
            try:
                return _extract_xlsx_text(p)
            except Exception:
                return None
    except Exception:
        return None
    return None
